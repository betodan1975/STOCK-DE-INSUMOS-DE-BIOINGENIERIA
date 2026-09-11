"""Migracion idempotente del Excel (hoja Stock) a SQLite + bootstrap del esquema.

Uso:
    cd api
    python migrate.py

- Crea / actualiza tablas (idempotente).
- Si la DB existe pero le faltan columnas nuevas (stock_minimo, usuario, etc.)
  las agrega via ALTER TABLE.
- Crea la tabla `usuarios` y siembra un admin/admin1234 si no hay ningun usuario.
- Si existe el Excel original, hace upsert por nombre.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from sqlalchemy import inspect, text

from auth import hash_password
from database import SessionLocal, engine
from models import Base, Insumo, Usuario

BASE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = BASE_DIR.parent
EXCEL_PATH = PROJECT_DIR / "Stock_Insumos_Dashboard.xlsx"
FOTOS_DIR = PROJECT_DIR / "fotos_insumos"


# -------------------- Schema migration helpers ----------------------------- #
def _ensure_columns() -> None:
    """ALTER TABLE idempotente para columnas agregadas en versiones nuevas."""
    insp = inspect(engine)

    def cols(table: str) -> set[str]:
        try:
            return {c["name"] for c in insp.get_columns(table)}
        except Exception:
            return set()

    statements: list[str] = []
    if "insumos" in insp.get_table_names():
        ins_cols = cols("insumos")
        if "stock_minimo" not in ins_cols:
            statements.append(
                "ALTER TABLE insumos ADD COLUMN stock_minimo INTEGER NOT NULL DEFAULT 1"
            )
    if "movimientos" in insp.get_table_names():
        mov_cols = cols("movimientos")
        if "usuario" not in mov_cols:
            statements.append("ALTER TABLE movimientos ADD COLUMN usuario VARCHAR")

    if statements:
        with engine.begin() as conn:
            for sql in statements:
                conn.execute(text(sql))
                print(f"  - {sql}")


def _seed_admin() -> None:
    db = SessionLocal()
    try:
        if db.query(Usuario).count() > 0:
            return
        admin = Usuario(
            username="admin",
            password_hash=hash_password("admin1234"),
            nombre="Administrador",
            rol="admin",
            activo=True,
        )
        db.add(admin)
        db.commit()
        print("=" * 60)
        print("USUARIO ADMIN INICIAL CREADO")
        print("  username: admin")
        print("  password: admin1234")
        print("  CAMBIALO desde la app o via POST /auth/users apenas puedas!")
        print("=" * 60)
    finally:
        db.close()


# Orden de columnas en la hoja Stock (fila 2 es encabezado):
# A Foto | B ITEM | C Stock | D Estado | E Equipo medico | F Marca |
# G Referencia/URL | H Ubicacion | I Prom. anual | J Solicitud compra |
# K Moneda | L Precio (original) | M Precio unit. valor | N Precio USD | O Precio ARS


def _clean_str(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _clean_int(v) -> int | None:
    if v is None or v == "":
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _clean_float(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _build_foto_index() -> dict[int, str]:
    """Devuelve {ordinal: filename} para todas las fotos tipo "NN_*.ext" en fotos_insumos/."""
    index: dict[int, str] = {}
    if not FOTOS_DIR.exists():
        return index
    pat = re.compile(r"^(\d{1,3})_.*\.(png|jpe?g|webp|gif)$", re.IGNORECASE)
    for f in FOTOS_DIR.iterdir():
        if not f.is_file():
            continue
        m = pat.match(f.name)
        if m:
            ordinal = int(m.group(1))
            if ordinal not in index:
                index[ordinal] = f.name
    return index


def run() -> None:
    # 1) Crear tablas que falten + ALTER TABLE para columnas nuevas + seed admin.
    Base.metadata.create_all(bind=engine)
    _ensure_columns()
    _seed_admin()

    if not EXCEL_PATH.exists():
        print(f"(!) No encuentro el Excel original ({EXCEL_PATH.name}); "
              "salteo la importacion inicial. Esto es esperable si ya migraste "
              "antes o estas laburando sobre datos cargados a mano.")
        return

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    if "Stock" not in wb.sheetnames:
        raise SystemExit("El Excel no tiene una hoja llamada 'Stock'.")
    ws = wb["Stock"]

    fotos = _build_foto_index()

    session = SessionLocal()
    created = updated = skipped = 0

    # La fila 1 es titulo decorativo, la 2 son encabezados, datos empiezan en la 3.
    ordinal = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or all(v is None for v in row):
            continue
        nombre = _clean_str(row[1])
        if not nombre:
            skipped += 1
            continue
        ordinal += 1

        fields = dict(
            nombre=nombre,
            stock=_clean_int(row[2]) or 0,
            estado=_clean_str(row[3]),
            equipo_medico=_clean_str(row[4]),
            marca=_clean_str(row[5]),
            referencia=_clean_str(row[6]),
            ubicacion=_clean_str(row[7]),
            prom_anual=_clean_int(row[8]),
            solicitud_compra=_clean_int(row[9]),
            moneda=_clean_str(row[10]),
            precio_original=_clean_str(row[11]),
            precio_unit_valor=_clean_float(row[12]),
            foto_filename=fotos.get(ordinal),
        )

        existing = session.query(Insumo).filter_by(nombre=nombre).one_or_none()
        if existing is None:
            session.add(Insumo(**fields))
            created += 1
        else:
            for k, v in fields.items():
                setattr(existing, k, v)
            updated += 1

    session.commit()
    session.close()

    print(f"Migracion OK -> creados: {created}, actualizados: {updated}, saltados: {skipped}")
    print(f"DB: {BASE_DIR / 'stock.db'}")


if __name__ == "__main__":
    run()
